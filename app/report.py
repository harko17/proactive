from __future__ import annotations

import re
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


class ExcelReportWriter:
    def __init__(self, output_path: str):
        self.output_path = Path(output_path)
        self.output_path.parent.mkdir(parents=True, exist_ok=True)

    def write(self, rows: list[dict]) -> str:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Dashboard Summary"

        headers = [
            "Dashboard",
            "Page",
            "Category",
            "Widget",
            "Visualization",
            "Account ID",
            "Metric Type",
            "Value",
            "Timestamp (UTC)",
            "Status",
            "Source",
        ]
        dashboard_name = next((row.get("dashboard_name") for row in rows if row.get("dashboard_name")), "New Relic Dashboard")
        dashboard_url = next((row.get("dashboard_url") for row in rows if row.get("dashboard_url")), "")
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        title_cell = sheet.cell(row=1, column=1, value=f"{dashboard_name} - Monitoring charts")
        title_cell.font = Font(name="Aptos Display", size=16, bold=True, color="FFFFFF")
        title_cell.fill = PatternFill("solid", fgColor="17365D")
        title_cell.alignment = Alignment(vertical="center")
        if dashboard_url.lower().startswith(("https://", "http://")):
            title_cell.hyperlink = dashboard_url
            title_cell.font = Font(name="Aptos Display", size=16, bold=True, color="FFFFFF", underline="single")
        sheet.row_dimensions[1].height = 28

        sheet.cell(row=2, column=1, value="Report generated (UTC)")
        sheet.cell(row=2, column=1).font = Font(bold=True, color="17365D")
        sheet.cell(row=2, column=2, value=datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC"))

        for column, header in enumerate(headers, start=1):
            cell = sheet.cell(row=4, column=column, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="4472C4")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for row_index, row in enumerate(rows, start=5):
            values = [
                row.get("dashboard_name", dashboard_name),
                row.get("page", ""),
                row.get("category", ""),
                row.get("metric_name", ""),
                row.get("visualization", ""),
                row.get("account_id", ""),
                row.get("metric_type", ""),
                self._excel_value(row.get("value")),
                row.get("timestamp", ""),
                row.get("status", "OK" if row.get("value") is not None else "No value"),
                row.get("source", ""),
            ]
            for column, value in enumerate(values, start=1):
                cell = sheet.cell(row=row_index, column=column, value=value)
                cell.alignment = Alignment(vertical="top", wrap_text=column in {1, 2, 3, 4, 9, 10})
                if column != 8:
                    cell.data_type = "s"
            if row_index % 2 == 0:
                for cell in sheet[row_index]:
                    cell.fill = PatternFill("solid", fgColor="F3F6FA")

        sheet.freeze_panes = "A5"
        sheet.auto_filter.ref = f"A4:{get_column_letter(len(headers))}{max(4, sheet.max_row)}"
        sheet.row_dimensions[4].height = 30
        widths = [34, 24, 30, 44, 18, 16, 18, 16, 26, 24, 18]
        for column, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(column)].width = width
        workbook.save(self.output_path)
        return str(self.output_path)

    @staticmethod
    def _excel_value(value):
        if isinstance(value, bool) or value is None:
            return value
        if isinstance(value, (int, float)):
            return value
        if isinstance(value, str):
            try:
                return float(value) if "." in value else int(value)
            except ValueError:
                return value
        return str(value)


def normalize_value(raw_value: str) -> str:
    value = raw_value.strip()
    value = value.replace(",", "")
    value = value.replace("%", "")
    value = value.replace(" ", "")

    if value.lower().endswith("k"):
        value = str(float(value[:-1]) * 1000)
    elif value.lower().endswith("m"):
        value = str(float(value[:-1]) * 1000000)

    if value.endswith(".0"):
        value = value[:-2]

    return value


def extract_metric_value(page_text: str, metric_rule: dict) -> str:
    pattern = metric_rule.get("regex", "")
    if not pattern:
        return ""

    match = re.search(pattern, page_text, re.IGNORECASE | re.DOTALL)
    if not match:
        return ""

    value = match.group(1).strip()
    return normalize_value(value)


def build_rows(page_text: str, metric_rules: list[dict]) -> list[dict]:
    rows: list[dict] = []
    timestamp = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")

    for rule in metric_rules:
        value = extract_metric_value(page_text, rule)
        if value != "":
            rows.append(
                {
                    "timestamp": timestamp,
                    "category": rule.get("category", ""),
                    "metric_name": rule.get("name", ""),
                    "metric_type": rule.get("metric_type", ""),
                    "value": value,
                    "unit": rule.get("unit", ""),
                    "source": "newrelic dashboard",
                }
            )

    return rows
