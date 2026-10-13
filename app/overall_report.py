from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

from app.market import _match_score


MARKET_COLUMNS = {
    "ZA": "za",
    "GLO DE": "glo-de",
    "GLO PL": "glo-pl",
    "Japan": "japan",
    "Korea": "korea",
    "B2B KR": "b2b-kr",
    "VELO PL": "velo-pl",
    "VUSE DE": "vuse-de",
    "GLO IT": "glo-it",
    "Serbia": "serbia",
    "VELO IT": "velo-it",
    "VUSE IT": "vuse-it",
    "Tetris FR": "tetris-fr",
    "Tetris GR": "tetris-gr",
}


def _cells(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _metric_key(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", value.casefold())


def _read_template(path: Path) -> tuple[list[str], list[list[str]]]:
    table_rows = [_cells(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip().startswith("|")]
    if not table_rows:
        raise ValueError(f"No Markdown table found in overall report template: {path}")

    headers = table_rows[0]
    expected_headers = ["Areas", *MARKET_COLUMNS]
    if headers != expected_headers:
        raise ValueError(
            f"Overall report headers in {path} must be {expected_headers!r}; found {headers!r}"
        )

    metric_rows = [row for row in table_rows[1:] if len(row) == len(headers) and row[0]]
    if not metric_rows:
        raise ValueError(f"No metric rows found in overall report template: {path}")
    return headers, metric_rows


def _read_market_values(path: Path) -> list[tuple[str, object]]:
    if not path.is_file():
        raise FileNotFoundError(f"Market report workbook is missing: {path}")
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        if "Dashboard Summary" not in workbook.sheetnames:
            raise ValueError(f"Market report workbook has no Dashboard Summary sheet: {path}")
        sheet = workbook["Dashboard Summary"]
        return [
            (str(row[1]).strip(), row[3])
            for row in sheet.iter_rows(min_row=2, values_only=True)
            if len(row) >= 4 and row[1] is not None and str(row[1]).strip()
        ]
    finally:
        workbook.close()


def _format_value(value: object) -> str:
    if value is None:
        return "-"
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float):
        return format(value, ".15g")
    return str(value).strip() or "-"


def _find_value(metric: str, report_rows: list[tuple[str, object]]) -> str:
    normalized_metric = _metric_key(metric)
    exact_matches = [value for label, value in report_rows if _metric_key(label) == normalized_metric]
    if exact_matches:
        return _format_value(exact_matches[0])

    matches = sorted(
        ((_match_score(label, metric), label, value) for label, value in report_rows),
        key=lambda match: (match[0], match[1]),
        reverse=True,
    )
    if not matches or matches[0][0] < 0.9:
        return "-"
    if len(matches) > 1 and matches[0][0] - matches[1][0] < 0.06:
        return "-"
    return _format_value(matches[0][2])


def _render_table(headers: list[str], rows: list[list[str]]) -> str:
    widths = [
        max(len(header), *(len(row[index]) for row in rows))
        for index, header in enumerate(headers)
    ]

    def border() -> str:
        return "+" + "+".join("-" * (width + 2) for width in widths) + "+"

    def render_row(row: list[str]) -> str:
        return "| " + " | ".join(value.ljust(widths[index]) for index, value in enumerate(row)) + " |"

    output = [border(), render_row(headers), border()]
    for row in rows:
        output.extend((render_row(row), border()))
    return "\n".join(output) + "\n"


def generate_overall_report(template_path: Path, reports_dir: Path, output_path: Path) -> str:
    headers, metric_rows = _read_template(template_path)
    values_by_market = {
        market: _read_market_values(reports_dir / f"{slug}.xlsx")
        for market, slug in MARKET_COLUMNS.items()
    }

    report_rows = []
    for template_row in metric_rows:
        metric = template_row[0]
        report_rows.append([
            metric,
            *[
                _find_value(metric, values_by_market[market])
                for market in MARKET_COLUMNS
            ],
        ])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(_render_table(headers, report_rows), encoding="utf-8")
    return str(output_path)
