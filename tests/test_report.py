import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.config import get_runtime_settings
from app.market import apply_report_layout, load_request, parse_report_layout, parse_graphql_entries
from app.newrelic_client import NewRelicDashboardClient
from app.report import ExcelReportWriter


PROJECT_ROOT = Path(__file__).resolve().parents[1]
MARKETS_ROOT = PROJECT_ROOT / "markets"
MARKET_DIRS = sorted(path for path in MARKETS_ROOT.iterdir() if path.is_dir())


def make_client(**settings):
    client = object.__new__(NewRelicDashboardClient)
    client.settings = {
        "markets_root": str(MARKETS_ROOT),
        "output_dir": str(PROJECT_ROOT / "output" / "test"),
        "api_key": "test-api-key",
        "graphql_endpoint": "https://api.newrelic.com/graphql",
        **settings,
    }
    return client


def load_response_fixture(path: Path):
    text = path.read_text(encoding="utf-8")
    decoder = json.JSONDecoder()
    for index, character in enumerate(text):
        if character != "{":
            continue
        try:
            value, _ = decoder.raw_decode(text[index:])
        except json.JSONDecodeError:
            continue
        if isinstance(value, dict) and isinstance(value.get("data"), dict):
            return value
    return None


def make_mock_response(entries):
    actor = {}
    for index, entry in enumerate(entries, start=1):
        if "TIMESERIES" in entry["query"].upper():
            results = [
                {"beginTimeSeconds": 1_791_686_399, "endTimeSeconds": 1_791_688_199, "count": 0},
                {"beginTimeSeconds": 1_791_688_199, "endTimeSeconds": 1_791_689_999, "count": index},
                {"beginTimeSeconds": 1_791_689_999, "endTimeSeconds": 1_791_691_799, "count": 0},
            ]
        else:
            results = [{"count": index}]
        actor[entry["alias"]] = {"nrql": {"results": results}}
    return {"data": {"actor": actor}}


@pytest.mark.parametrize("market_dir", MARKET_DIRS, ids=lambda path: path.name)
def test_every_market_request_has_unique_aliases_and_a_complete_report_row_set(market_dir):
    _, entries = load_request(market_dir / "CURLS.md")
    layout = parse_report_layout(market_dir / "EXCEL_REPORT.md", market_dir.name)
    rows = apply_report_layout(entries, layout)

    assert entries
    assert len({entry["alias"] for entry in entries}) == len(entries)
    assert len(rows) >= len(entries)
    assert {row["alias"] for row in rows if row.get("alias")} == {
        entry["alias"] for entry in entries
    }


@pytest.mark.parametrize(
    ("market_name", "expected_title", "expected_metrics"),
    [
        ("tetris-fr", "Tetris FR", 24),
        ("tetris-gr", "Tetris GR", 12),
    ],
)
def test_tetris_report_layout_matches_every_query(market_name, expected_title, expected_metrics):
    market_dir = MARKETS_ROOT / market_name
    _, entries = load_request(market_dir / "CURLS.md")
    layout = parse_report_layout(market_dir / "EXCEL_REPORT.md", market_name)
    rows = apply_report_layout(entries, layout)

    assert layout["title"] == expected_title
    assert len(entries) == expected_metrics
    assert len(rows) == expected_metrics
    assert all(row["mapping_status"] == "matched" for row in rows)


@pytest.mark.parametrize("market_dir", MARKET_DIRS, ids=lambda path: path.name)
def test_saved_response_fixture_matches_every_market_query_alias(market_dir):
    response = load_response_fixture(market_dir / "response.md")
    if response is None:
        pytest.skip(f"{market_dir.name} has no populated market-local response fixture")

    _, entries = load_request(market_dir / "CURLS.md")
    actor = response["data"].get("actor")
    assert isinstance(actor, dict)
    expected_aliases = {entry["alias"] for entry in entries}

    def find_aliases(value):
        if isinstance(value, dict):
            return (set(value) & expected_aliases).union(
                *(find_aliases(child) for child in value.values())
            )
        if isinstance(value, list):
            return set().union(*(find_aliases(child) for child in value))
        return set()

    assert find_aliases(actor) == expected_aliases


def test_account_ids_are_taken_from_each_alias_query():
    _, japan_entries = load_request(MARKETS_ROOT / "japan" / "CURLS.md")
    japan = {entry["alias"]: entry for entry in japan_entries}
    assert japan["normalOrderCreation"]["account_id"] == "4198216"
    assert japan["productUpdatePricebookError"]["account_id"] == "6263774"
    assert japan["customerCreationYetToSyncSitecoreCrm"]["account_id"] == "6263774"

    _, glo_it_entries = load_request(MARKETS_ROOT / "glo-it" / "CURLS.md")
    assert {entry["account_id"] for entry in glo_it_entries} == {"6358718"}

    _, b2b_entries = load_request(MARKETS_ROOT / "b2b-kr" / "CURLS.md")
    assert {entry["account_id"] for entry in b2b_entries} == {"3939452"}


def test_short_korea_customer_create_layout_label_maps_to_trend_alias():
    market_dir = MARKETS_ROOT / "korea"
    _, entries = load_request(market_dir / "CURLS.md")
    layout = parse_report_layout(market_dir / "EXCEL_REPORT.md", market_dir.name)
    rows = apply_report_layout(entries, layout)

    customer_create = next(row for row in rows if row["metric_name"] == "Customer create")
    assert customer_create["alias"] == "customercreationsynctoCRMTrend"


def test_graphql_parser_supports_nested_account_and_nrql_accounts_forms():
    query = """
    query {
      actor {
        first: account(id: 4198216) {
          nrql(query: "FROM Log SELECT count(*) SINCE 24 HOURS AGO") { results }
        }
        second: nrql(accounts: [6358718], query: "FROM Log SELECT count(*) SINCE 24 HOURS AGO") { results }
      }
    }
    """

    assert parse_graphql_entries(query) == [
        {
            "alias": "first",
            "account_id": "4198216",
            "query": "FROM Log SELECT count(*) SINCE 24 HOURS AGO",
        },
        {
            "alias": "second",
            "account_id": "6358718",
            "query": "FROM Log SELECT count(*) SINCE 24 HOURS AGO",
        },
    ]


def test_latest_nonzero_time_series_bucket_wins_even_if_last_is_zero():
    buckets = [
        {"beginTimeSeconds": 1_791_686_399, "endTimeSeconds": 1_791_688_199, "count": 0},
        {"beginTimeSeconds": 1_791_688_199, "endTimeSeconds": 1_791_689_999, "count": 0},
        {"beginTimeSeconds": 1_791_689_999, "endTimeSeconds": 1_791_691_799, "count": 1},
        {"beginTimeSeconds": 1_791_691_799, "endTimeSeconds": 1_791_693_599, "count": 0},
    ]

    value, timestamp, status = NewRelicDashboardClient._extract_report_value(buckets)

    assert value == 1
    assert timestamp == "2026-10-11 04:09:59 UTC"
    assert status == "OK"


def test_count_message_series_and_all_zero_series_are_numeric():
    buckets = [
        {"beginTimeSeconds": 1_791_686_399, "endTimeSeconds": 1_791_688_199, "count.message": 0},
        {"beginTimeSeconds": 1_791_688_199, "endTimeSeconds": 1_791_689_999, "count.message": 2},
        {"beginTimeSeconds": 1_791_689_999, "endTimeSeconds": 1_791_691_799, "count.message": 0},
    ]
    assert NewRelicDashboardClient._extract_report_value(buckets) == (
        2,
        "2026-10-11 03:39:59 UTC",
        "OK",
    )
    assert NewRelicDashboardClient._extract_report_value(
        [{"beginTimeSeconds": 1_791_686_399, "endTimeSeconds": 1_791_688_199, "count": 0}]
    ) == (0, "2026-10-11 03:09:59 UTC", "OK")


def test_case_insensitive_nrql_result_alias_is_extracted():
    buckets = [
        {"beginTimeSeconds": 1_791_686_399, "endTimeSeconds": 1_791_688_199, "Counts": 0},
        {"beginTimeSeconds": 1_791_688_199, "endTimeSeconds": 1_791_689_999, "Counts": 40},
        {"beginTimeSeconds": 1_791_689_999, "endTimeSeconds": 1_791_691_799, "Counts": 0},
    ]
    value, _, status = NewRelicDashboardClient._extract_report_value(buckets)
    assert (value, status) == (40, "OK")


def test_graphql_alias_errors_do_not_discard_successful_results():
    entries = [{"alias": "good"}, {"alias": "bad"}]
    response = {
        "data": {"actor": {"good": {"nrql": {"results": [{"count": 4}]}}}},
        "errors": [{"message": "Account access denied", "path": ["actor", "bad", "nrql"]}],
    }

    assert NewRelicDashboardClient._extract_graphql_batch_results(response, entries) == {
        "good": {"results": [{"count": 4}]},
        "bad": {"error": "Account access denied"},
    }


def test_graphql_batch_results_support_aliases_inside_unaliased_account():
    entries = [{"alias": "good"}, {"alias": "bad"}]
    response = {
        "data": {
            "actor": {
                "account": {
                    "good": {"results": [{"count": 4}]},
                }
            }
        },
        "errors": [
            {"message": "NRQL query failed", "path": ["actor", "account", "bad", "nrql"]}
        ],
    }

    assert NewRelicDashboardClient._extract_graphql_batch_results(response, entries) == {
        "good": {"results": [{"count": 4}]},
        "bad": {"error": "NRQL query failed"},
    }


def test_tetris_market_requires_its_dedicated_api_key():
    client = make_client(tetris_api_key="")
    payload, entries = load_request(MARKETS_ROOT / "tetris-fr" / "CURLS.md")

    with pytest.raises(ValueError, match="TETRIS_API_KEY"):
        client._send_graphql_payload(
            payload,
            entries,
            api_key_name="TETRIS_API_KEY",
        )


@pytest.mark.parametrize(
    "market_dir",
    [path for path in MARKET_DIRS if path.name != "japan"],
    ids=lambda path: path.name,
)
def test_saved_market_responses_yield_values_or_explicit_empty_results(market_dir):
    response = load_response_fixture(market_dir / "response.md")
    assert response is not None
    _, entries = load_request(market_dir / "CURLS.md")
    results = NewRelicDashboardClient._extract_graphql_batch_results(response, entries)

    assert set(results) == {entry["alias"] for entry in entries}
    for alias_result in results.values():
        assert "error" in alias_result or isinstance(alias_result.get("results"), list)
        if "results" in alias_result:
            _, _, status = NewRelicDashboardClient._extract_report_value(alias_result["results"])
            assert status in {"OK", "No numeric result"}


def test_graphql_response_maps_actual_values_for_every_market():
    client = make_client()
    for market_dir in MARKET_DIRS:
        _, query_entries = load_request(market_dir / "CURLS.md")
        layout = parse_report_layout(market_dir / "EXCEL_REPORT.md", market_dir.name)
        entries = apply_report_layout(query_entries, layout)
        rows = client._rows_from_graphql_response(
            make_mock_response(query_entries),
            entries,
            layout,
        )

        assert len(rows) == len(entries)
        assert all(row["value"] is not None for row in rows if row["status"] == "OK")
        assert all(row["status"] == "No query configured" for row in rows if not row["account_id"])


def test_report_writer_preserves_repeated_categories_and_title_column(tmp_path):
    rows = [
        {"category": "Order Fulfillment", "metric_name": "Pending", "metric_type": "count", "value": 0},
        {"category": "Order Fulfillment", "metric_name": "Shipment", "metric_type": "trend", "value": 1},
    ]
    saved_path = ExcelReportWriter(
        str(tmp_path / "market.xlsx"),
        report_title="VUSE IT",
        title_column=2,
        merge_categories=False,
    ).write(rows)
    sheet = load_workbook(saved_path, data_only=True)["Dashboard Summary"]

    assert sheet["B1"].value == "VUSE IT - Monitoring charts"
    assert sheet["A2"].value == "Order Fulfillment"
    assert sheet["A3"].value == "Order Fulfillment"
    assert sheet["D2"].value == 0
    assert sheet["D3"].value == 1
    assert not sheet.merged_cells.ranges


def test_run_generates_market_workbooks_from_mocked_api_responses(monkeypatch, tmp_path):
    client = make_client(output_dir=str(tmp_path), tetris_api_key="test-tetris-api-key")
    credentials_by_alias = {}

    def mock_send(self, payload, entries, api_key=None, api_key_name="NEW_RELIC_API_KEY"):
        credentials_by_alias[entries[0]["alias"]] = (api_key, api_key_name)
        return make_mock_response(entries)

    monkeypatch.setattr(
        NewRelicDashboardClient,
        "_send_graphql_payload",
        mock_send,
    )

    report_paths = client.run()
    assert len(report_paths) == len(MARKET_DIRS)
    assert {Path(path).stem for path in report_paths} == {path.name for path in MARKET_DIRS}
    for path in report_paths:
        assert Path(path).is_file()
        workbook = load_workbook(path, data_only=True)
        assert workbook["Dashboard Summary"].max_row > 1
    assert credentials_by_alias["stuckinpendingstatus"] == (
        "test-tetris-api-key",
        "TETRIS_API_KEY",
    )
    assert credentials_by_alias["stuckInPendingStatus"] == (
        "test-tetris-api-key",
        "TETRIS_API_KEY",
    )
    assert credentials_by_alias["normalOrderCreation"] == (
        "test-api-key",
        "NEW_RELIC_API_KEY",
    )


def test_api_post_sends_exact_payload_and_postman_headers(monkeypatch):
    client = make_client()
    payload, entries = load_request(MARKETS_ROOT / "japan" / "CURLS.md")
    captured = {}

    class MockResponse:
        status_code = 200
        headers = {}
        text = ""

        @staticmethod
        def json():
            return make_mock_response(entries)

    def fake_post(url, headers, json, timeout):
        captured.update(url=url, headers=headers, payload=json, timeout=timeout)
        return MockResponse()

    monkeypatch.setattr("app.newrelic_client.requests.post", fake_post)
    response = client._send_graphql_payload(payload, entries)

    assert captured["url"] == "https://api.newrelic.com/graphql"
    assert captured["headers"] == {
        "API-Key": "test-api-key",
        "Content-Type": "application/json",
    }
    assert captured["payload"] == payload
    assert response["data"]["actor"]["normalOrderCreation"]["nrql"]["results"]


def test_runtime_settings_require_only_api_key(monkeypatch):
    monkeypatch.delenv("NEW_RELIC_API_KEY", raising=False)
    with pytest.raises(ValueError, match="NEW_RELIC_API_KEY"):
        get_runtime_settings()

    monkeypatch.setenv("NEW_RELIC_API_KEY", "test-key")
    settings = get_runtime_settings()

    assert settings["api_key"] == "test-key"
    assert settings["markets_root"].endswith("markets")
    assert settings["output_dir"].endswith("output")
